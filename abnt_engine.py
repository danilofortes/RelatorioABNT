# -*- coding: utf-8 -*-
"""Motor de formatacao ABNT para geracao de relatorios em DOCX."""

import os

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT_NAME = "Times New Roman"
FONT_SIZE_NORMAL = Pt(12)
FONT_SIZE_CODE = Pt(10)
FONT_SIZE_TITLE = Pt(12)
SUMARIO_TAB_POS = Cm(16)
# NBR 14724: espaco de 1,5 entre linhas com fonte 12 pt = 12 * 1.5 = 18 pt
ESPACO_1_5_LINHAS = Pt(18)


def set_font(run, name=FONT_NAME, size=FONT_SIZE_NORMAL, bold=False, italic=False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)


def add_paragraph_abnt(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                       indent=True, bold=False, size=FONT_SIZE_NORMAL,
                       space_after=Pt(0), line_spacing=1.5):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    if indent:
        pf.first_line_indent = Cm(1.25)
    pf.line_spacing = line_spacing
    pf.space_after = space_after
    pf.space_before = Pt(0)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p


def add_heading_abnt(doc, text, level=1, space_before=None, space_after=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    if level == 1:
        pf.space_before = space_before if space_before is not None else Pt(0)
        pf.space_after = space_after if space_after is not None else ESPACO_1_5_LINHAS
    else:
        pf.space_before = space_before if space_before is not None else ESPACO_1_5_LINHAS
        pf.space_after = space_after if space_after is not None else ESPACO_1_5_LINHAS
    run = p.add_run(text.upper() if level == 1 else text)
    set_font(run, size=FONT_SIZE_TITLE, bold=True)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._element.append(br)


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(6)
    pf.space_before = Pt(6)
    pf.left_indent = Cm(2.0)
    pf.right_indent = Cm(0.5)
    run = p.add_run(code_text)
    set_font(run, name="Consolas", size=FONT_SIZE_CODE)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)


def add_caption(doc, text, italic=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(6)
    pf.space_before = Pt(6)
    run = p.add_run(text)
    set_font(run, size=Pt(10), italic=italic)


def add_source(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(12)
    run = p.add_run(text)
    set_font(run, size=Pt(10))


def resolver_arquivo(arquivo, projeto_dir):
    if os.path.isabs(arquivo):
        caminho = arquivo
    else:
        caminho = os.path.join(projeto_dir, arquivo)
    if not os.path.isfile(caminho):
        raise FileNotFoundError(f"Arquivo de imagem nao encontrado: {caminho}")
    return caminho


def add_image_block(doc, image_path, largura_cm=14, altura_cm=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    run = p.add_run()
    if altura_cm:
        run.add_picture(image_path, width=Cm(largura_cm), height=Cm(altura_cm))
    else:
        run.add_picture(image_path, width=Cm(largura_cm))


def add_table_abnt(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_font(run, bold=True, size=Pt(11))
    for r, row in enumerate(rows):
        cells = table.rows[r + 1].cells
        for c, val in enumerate(row):
            cells[c].text = ""
            p = cells[c].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            set_font(run, size=Pt(11))
    return table


def configure_default_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE_NORMAL
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:cs"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)


def configure_margins(doc):
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.left_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.right_margin = Cm(2)
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)


def montar_capa(doc, meta):
    for txt in meta.get("instituicao_linhas", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(txt)
        set_font(run, size=Pt(12), bold=True)

    for _ in range(meta.get("espacos_antes_autores", 6)):
        doc.add_paragraph()

    for autor in meta.get("autores", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(autor)
        set_font(run, size=Pt(12))

    if meta.get("grupo"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(meta["grupo"])
        set_font(run, size=Pt(12))

    for _ in range(meta.get("espacos_antes_titulo", 5)):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(meta["titulo"])
    set_font(run, size=Pt(14), bold=True)

    if meta.get("subtitulo"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(meta["subtitulo"])
        set_font(run, size=Pt(12), italic=True)

    for _ in range(meta.get("espacos_antes_local", 8)):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(meta.get("cidade", "Campinas"))
    set_font(run, size=Pt(12))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(meta.get("ano", "2026")))
    set_font(run, size=Pt(12))

    add_page_break(doc)


def montar_folha_rosto(doc, meta):
    for autor in meta.get("autores", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(autor)
        set_font(run, size=Pt(12))

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(meta["titulo"])
    set_font(run, size=Pt(14), bold=True)

    if meta.get("subtitulo"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(meta["subtitulo"])
        set_font(run, size=Pt(12), italic=True)

    for _ in range(3):
        doc.add_paragraph()

    if meta.get("nota_folha_rosto"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Cm(8.0)
        pf.line_spacing = 1.0
        run = p.add_run(meta["nota_folha_rosto"])
        set_font(run, size=Pt(10))

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(meta.get("cidade", "Campinas"))
    set_font(run, size=Pt(12))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(meta.get("ano", "2026")))
    set_font(run, size=Pt(12))

    add_page_break(doc)


def add_sumario_item(doc, titulo, pagina, nivel=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    is_secao = nivel == 1
    recuo = Cm(0) if is_secao else Cm(0.75)
    pf.left_indent = recuo
    tab_pos = SUMARIO_TAB_POS - recuo
    pf.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    run = p.add_run(titulo)
    set_font(run, size=Pt(12), bold=is_secao)
    p.add_run("\t")
    run_pag = p.add_run(str(pagina))
    set_font(run_pag, size=Pt(12), bold=is_secao)


def montar_sumario(doc, itens):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("SUMÁRIO")
    set_font(run, size=Pt(12), bold=True)

    for item in itens:
        if isinstance(item, dict):
            titulo = item["titulo"]
            pagina = item["pagina"]
            nivel = item.get("nivel", 1)
        else:
            titulo, pagina = item[0], item[1]
            nivel = 1 if titulo[:1].isdigit() and ". " not in titulo.split(" ", 1)[0] else 2
        add_sumario_item(doc, titulo, pagina, nivel=nivel)

    add_page_break(doc)


def render_bloco_figura(doc, bloco, projeto_dir):
    if bloco.get("paragrafo_antes"):
        add_paragraph_abnt(doc, bloco["paragrafo_antes"])
    add_caption(doc, bloco["legenda"])

    if bloco.get("arquivo"):
        caminho = resolver_arquivo(bloco["arquivo"], projeto_dir)
        add_image_block(
            doc,
            caminho,
            largura_cm=bloco.get("largura_cm", 14),
            altura_cm=bloco.get("altura_cm"),
        )
    elif bloco.get("conteudo"):
        add_code_block(doc, bloco["conteudo"])
    else:
        raise ValueError(
            "Bloco 'figura' precisa de 'arquivo' (imagem) ou 'conteudo' (texto)."
        )

    if bloco.get("fonte"):
        add_source(doc, bloco["fonte"])


def render_bloco(doc, bloco, projeto_dir="."):
    tipo = bloco["tipo"]

    if tipo == "paragrafo":
        add_paragraph_abnt(doc, bloco["texto"])
        return

    if tipo == "quadro":
        add_caption(doc, bloco["legenda"])
        add_table_abnt(doc, bloco["cabecalhos"], bloco["linhas"])
        if bloco.get("fonte"):
            add_source(doc, bloco["fonte"])
        return

    if tipo == "codigo":
        if bloco.get("paragrafo_antes"):
            add_paragraph_abnt(doc, bloco["paragrafo_antes"])
        add_caption(doc, bloco["legenda"])
        add_code_block(doc, bloco["codigo"])
        if bloco.get("fonte"):
            add_source(doc, bloco["fonte"])
        return

    if tipo in ("figura", "imagem"):
        render_bloco_figura(doc, bloco, projeto_dir)
        return

    if tipo == "subsecao":
        add_heading_abnt(doc, bloco["titulo"], level=bloco.get("nivel", 2))
        for sub in bloco.get("blocos", []):
            render_bloco(doc, sub, projeto_dir=projeto_dir)
        return

    raise ValueError(f"Tipo de bloco desconhecido: {tipo}")


def render_secao(doc, secao, projeto_dir=".", espaco_antes=False):
    space_before = ESPACO_1_5_LINHAS if espaco_antes else None
    add_heading_abnt(
        doc,
        secao["titulo"],
        level=secao.get("nivel", 1),
        space_before=space_before,
    )
    for bloco in secao.get("blocos", []):
        render_bloco(doc, bloco, projeto_dir=projeto_dir)
    if secao.get("quebra_pagina"):
        add_page_break(doc)


def montar_referencias(doc, referencias, titulo="REFERÊNCIAS"):
    add_heading_abnt(doc, titulo)
    for ref in referencias:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = ESPACO_1_5_LINHAS
        pf.first_line_indent = Cm(0)
        run = p.add_run(ref)
        set_font(run, size=Pt(12))


def gerar_documento(
    meta,
    sumario,
    secoes,
    consideracoes,
    referencias,
    projeto_dir=".",
    titulos_finais=None,
):
    titulos = titulos_finais or {}
    titulo_consideracoes = titulos.get("consideracoes", "CONSIDERAÇÕES FINAIS")
    titulo_referencias = titulos.get("referencias", "REFERÊNCIAS")

    doc = Document()
    configure_default_style(doc)
    configure_margins(doc)

    montar_capa(doc, meta)
    montar_folha_rosto(doc, meta)
    montar_sumario(doc, sumario)

    primarias_nova_pagina = meta.get("secoes_primarias_nova_pagina", True)
    primeira_primaria = True
    for secao in secoes:
        nivel = secao.get("nivel", 1)
        espaco_antes = False
        if nivel == 1:
            if primarias_nova_pagina and not primeira_primaria:
                add_page_break(doc)
            elif not primarias_nova_pagina and not primeira_primaria:
                espaco_antes = True
            primeira_primaria = False
        render_secao(doc, secao, projeto_dir=projeto_dir, espaco_antes=espaco_antes)

    add_page_break(doc)
    add_heading_abnt(doc, titulo_consideracoes)
    for paragrafo in consideracoes:
        add_paragraph_abnt(doc, paragrafo)
    add_page_break(doc)

    montar_referencias(doc, referencias, titulo=titulo_referencias)
    return doc
